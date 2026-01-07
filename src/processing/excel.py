import io
import os
from typing import Any, Dict

import pandas as pd


# Export output JSON to a simple Excel file.
def _output_to_excel_bytes(output: Dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        if isinstance(output, dict) and isinstance(output.get("agreements"), list):
            agreements = output.get("agreements", [])
            if agreements:
                df = pd.json_normalize(agreements)
            else:
                df = pd.DataFrame()
            df.to_excel(writer, sheet_name="agreements", index=False)
            metadata = output.get("metadata", {})
            if metadata:
                pd.DataFrame([metadata]).to_excel(
                    writer, sheet_name="metadata", index=False
                )
        else:
            df = pd.json_normalize(output)
            df.to_excel(writer, sheet_name="output", index=False)
    return buffer.getvalue()


# Write a DataFrame into an openpyxl worksheet.
def _write_dataframe_to_sheet(worksheet: Any, df: pd.DataFrame) -> None:
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            worksheet.cell(row=row_idx, column=col_idx, value=value)


# Export output JSON into the template workbook with new sheets.
def _output_to_template_excel_bytes(output: Dict[str, Any]) -> bytes:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to build template output.") from exc

    template_path = os.path.join(
        os.path.dirname(__file__),
        "..",
        "..",
        "data",
        "templates",
        "template.xlsx",
    )
    workbook = load_workbook(template_path)

    for sheet_name in ("agreements", "metadata"):
        if sheet_name in workbook.sheetnames:
            workbook.remove(workbook[sheet_name])

    agreements = output.get("agreements") if isinstance(output, dict) else None
    if isinstance(agreements, list):
        ws = workbook.create_sheet("agreements")
        df = pd.json_normalize(agreements)
        if not df.empty:
            _write_dataframe_to_sheet(ws, df)

    metadata = output.get("metadata") if isinstance(output, dict) else None
    if isinstance(metadata, dict) and metadata:
        ws = workbook.create_sheet("metadata")
        df = pd.DataFrame([metadata])
        _write_dataframe_to_sheet(ws, df)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
